import io
import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

class DataExporter:
    @staticmethod
    def get_timestamped_filename(product_name: str, file_type: str) -> str:
        """
        Generates safe filenames like: Smart_Watch_2026-07-09.csv
        """
        # Sanitize product name for file systems
        safe_name = "".join(c if c.isalnum() else "_" for c in product_name).strip("_")
        # Keep it clean
        safe_name = "_".join(filter(None, safe_name.split("_")))
        date_str = datetime.date.today().strftime("%Y-%m-%d")
        return f"{safe_name}_{date_str}.{file_type}"

    @staticmethod
    def export_to_csv(df: pd.DataFrame) -> bytes:
        """
        Converts the competitor DataFrame into an encoded CSV byte-stream.
        """
        if df.empty:
            return b""
            
        # Create CSV in memory
        csv_buffer = io.StringIO()
        df.to_csv(csv_buffer, index=False)
        return csv_buffer.getvalue().encode("utf-8")

    @staticmethod
    def export_to_excel(df: pd.DataFrame, ai_listing: dict = None) -> bytes:
        """
        Generates a professionally-formatted multi-sheet Excel workbook in-memory.
        - Sheet 1: Competitor Analysis Table
        - Sheet 2: AI-Generated Listing Copy (if available)
        """
        output = io.BytesIO()
        wb = Workbook()
        
        # Style Definitions
        navy_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        light_gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        white_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        bold_font = Font(name="Arial", size=11, bold=True)
        regular_font = Font(name="Arial", size=10)
        title_font = Font(name="Arial", size=14, bold=True, color="1F497D")
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        
        # -------------------------------------------------------------
        # SHEET 1: Competitor Products
        # -------------------------------------------------------------
        ws_comp = wb.active
        ws_comp.title = "Competitor Market Data"
        ws_comp.views.sheetView[0].showGridLines = True
        
        # Add a title row
        ws_comp.append(["Competitor Market Pricing and Performance Analysis"])
        ws_comp.cell(row=1, column=1).font = title_font
        ws_comp.row_dimensions[1].height = 25
        ws_comp.append([])  # blank row
        
        if not df.empty:
            # Let's clean up dataframe for sheet view
            export_df = df.copy()
            # Rename columns to look professional
            column_mapping = {
                "name": "Product Name",
                "price": "Price (Normalized)",
                "rating": "Rating",
                "reviews": "Review Count",
                "availability": "Stock Availability",
                "website": "E-Commerce Website",
                "url": "Product URL"
            }
            # Filter and reorder
            available_cols = [col for col in column_mapping.keys() if col in export_df.columns]
            export_df = export_df[available_cols].rename(columns=column_mapping)
            
            # Write headers
            headers = list(export_df.columns)
            ws_comp.append(headers)
            header_row_idx = 3
            ws_comp.row_dimensions[header_row_idx].height = 22
            
            for col_idx, h in enumerate(headers, 1):
                cell = ws_comp.cell(row=header_row_idx, column=col_idx)
                cell.fill = navy_fill
                cell.font = white_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            # Write data rows
            for r_idx, row in enumerate(dataframe_to_rows(export_df, index=False, header=False), start=4):
                ws_comp.append(row)
                ws_comp.row_dimensions[r_idx].height = 20
                for col_idx in range(1, len(row) + 1):
                    cell = ws_comp.cell(row=r_idx, column=col_idx)
                    cell.font = regular_font
                    cell.border = thin_border
                    
                    # Align price and numbers to right
                    if col_idx in [2, 4]:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                        
                    # Give links a hyperlinked look if it is the URL column
                    if headers[col_idx-1] == "Product URL" and cell.value:
                        cell.font = Font(name="Arial", size=10, underline="single", color="0563C1")
        else:
            ws_comp.append(["No competitor market data scraped."])
            ws_comp.cell(row=3, column=1).font = regular_font

        # Auto-fit columns
        for col in ws_comp.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            col_letter = col[0].column_letter
            ws_comp.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # -------------------------------------------------------------
        # SHEET 2: AI-Generated Listing Copy
        # -------------------------------------------------------------
        if ai_listing:
            ws_ai = wb.create_sheet(title="AI Product Copy")
            ws_ai.views.sheetView[0].showGridLines = True
            
            ws_ai.append(["AI-Generated SEO Optimized E-Commerce Listing"])
            ws_ai.cell(row=1, column=1).font = title_font
            ws_ai.row_dimensions[1].height = 25
            ws_ai.append([]) # blank
            
            fields_to_write = [
                ("SEO Optimized Title", ai_listing.get("seo_title", "")),
                ("Short Description / Pitch", ai_listing.get("short_description", "")),
                ("Long Meta Description (SEO)", ai_listing.get("meta_description", "")),
                ("Call To Action", ai_listing.get("call_to_action", "")),
                ("Detailed Product Description", ai_listing.get("long_description", "")),
            ]
            
            current_row = 3
            for title, content in fields_to_write:
                # Section Header
                ws_ai.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
                ws_ai.cell(row=current_row, column=1, value=title).font = bold_font
                ws_ai.cell(row=current_row, column=1).fill = light_gray_fill
                ws_ai.row_dimensions[current_row].height = 20
                current_row += 1
                
                # Section Content
                ws_ai.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
                cell = ws_ai.cell(row=current_row, column=1, value=content)
                cell.font = regular_font
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                
                # Dynamic row height calculation based on content length
                approx_lines = max(1, len(str(content)) // 90)
                ws_ai.row_dimensions[current_row].height = max(20, approx_lines * 15)
                current_row += 2 # leave an extra blank row after
                
            # Bullets Section
            ws_ai.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
            ws_ai.cell(row=current_row, column=1, value="AI-Generated Bullet Points").font = bold_font
            ws_ai.cell(row=current_row, column=1).fill = light_gray_fill
            ws_ai.row_dimensions[current_row].height = 20
            current_row += 1
            
            bullets = ai_listing.get("bullets", [])
            for bullet in bullets:
                ws_ai.cell(row=current_row, column=1, value="•").alignment = Alignment(horizontal="center", vertical="top")
                cell = ws_ai.cell(row=current_row, column=2, value=bullet)
                cell.font = regular_font
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                approx_lines = max(1, len(str(bullet)) // 80)
                ws_ai.row_dimensions[current_row].height = max(18, approx_lines * 14)
                current_row += 1
                
            current_row += 1
            
            # Features Section
            ws_ai.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
            ws_ai.cell(row=current_row, column=1, value="Core Features").font = bold_font
            ws_ai.cell(row=current_row, column=1).fill = light_gray_fill
            ws_ai.row_dimensions[current_row].height = 20
            current_row += 1
            
            features = ai_listing.get("features", [])
            for feat in features:
                ws_ai.cell(row=current_row, column=1, value="✓").alignment = Alignment(horizontal="center", vertical="top")
                cell = ws_ai.cell(row=current_row, column=2, value=feat)
                cell.font = regular_font
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws_ai.row_dimensions[current_row].height = 18
                current_row += 1
                
            current_row += 1
            
            # Specifications Section
            ws_ai.cell(row=current_row, column=1, value="Technical Specification").font = bold_font
            ws_ai.cell(row=current_row, column=1).fill = light_gray_fill
            ws_ai.cell(row=current_row, column=2, value="Value").font = bold_font
            ws_ai.cell(row=current_row, column=2).fill = light_gray_fill
            ws_ai.row_dimensions[current_row].height = 20
            current_row += 1
            
            specs = ai_listing.get("specifications", {})
            for key, val in specs.items():
                cell_k = ws_ai.cell(row=current_row, column=1, value=key)
                cell_v = ws_ai.cell(row=current_row, column=2, value=val)
                cell_k.font = regular_font
                cell_v.font = regular_font
                cell_k.border = thin_border
                cell_v.border = thin_border
                ws_ai.row_dimensions[current_row].height = 18
                current_row += 1
                
            current_row += 1
            
            # Keywords and Selling Points (Side by Side or stacked)
            ws_ai.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
            ws_ai.cell(row=current_row, column=1, value="Search Backend SEO Keywords").font = bold_font
            ws_ai.cell(row=current_row, column=1).fill = light_gray_fill
            ws_ai.row_dimensions[current_row].height = 20
            current_row += 1
            
            kw_str = ", ".join(ai_listing.get("keywords", []))
            ws_ai.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
            cell = ws_ai.cell(row=current_row, column=1, value=kw_str)
            cell.font = regular_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws_ai.row_dimensions[current_row].height = 30
            
            # Configure columns
            ws_ai.column_dimensions["A"].width = 28
            ws_ai.column_dimensions["B"].width = 65

        # Save workbook to memory buffer
        wb.save(output)
        return output.getvalue()
